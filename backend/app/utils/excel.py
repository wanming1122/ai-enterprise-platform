"""Excel 读写工具（openpyxl）：导出文件生成与导入文件解析。"""
from io import BytesIO

from openpyxl import Workbook, load_workbook


def export_workbook(headers: list[str], rows: list[list]) -> bytes:
    """按表头与数据行生成 xlsx 字节流。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_workbook(content: bytes, headers: list[str]) -> list[dict]:
    """解析 xlsx：按给定表头列名返回 [{表头: 值}]，跳过空行，单元格空值归一为空串。"""
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []

    header_row = [str(h).strip() if h is not None else "" for h in rows[0]]
    col_index = {name: i for i, name in enumerate(header_row)}

    result: list[dict] = []
    for raw in rows[1:]:
        if all(c is None or str(c).strip() == "" for c in raw):
            continue
        row: dict = {}
        for name in headers:
            i = col_index.get(name)
            val = raw[i] if i is not None and i < len(raw) else None
            row[name] = str(val).strip() if val is not None else ""
        result.append(row)
    return result
